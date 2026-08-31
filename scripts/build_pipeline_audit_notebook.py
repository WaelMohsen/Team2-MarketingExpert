"""Generate the executable Sample 2 pipeline audit notebook."""

from pathlib import Path
from textwrap import dedent

import nbformat as nbf


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "src_2_pipeline_step_by_step.ipynb"


def markdown(text: str):
    return nbf.v4.new_markdown_cell(dedent(text).strip())


def code(text: str):
    return nbf.v4.new_code_cell(dedent(text).strip())


cells = [
    markdown(
        """
        # Sample 2 Completed-Cycle Pipeline: Step-by-Step Audit

        This notebook executes the current `src_2` pipeline one stage at a time and
        records the output of every completed stage in one Excel workbook.

        **Business objective:** use the completed batch to score campaigns and their
        adsets, ads, creatives, and audiences from observed WhatsApp business outcomes,
        then create an illustrative next-cycle budget scenario. Meta metrics explain
        delivery behavior; they do not replace delivered-order outcomes.

        The notebook keeps Meta media, WhatsApp conversations, and order lines at their
        natural grains. It aggregates them separately before joining scorecards, which
        prevents duplicated spend or duplicated orders.

        The LLM section is disabled by default to avoid unplanned API calls. Set
        `RUN_LLM = True` in the setup cell to execute the 12 campaign analyses, one
        portfolio synthesis, and one stakeholder report call.
        """
    ),
    markdown(
        """
        ## Pipeline map

        1. Read the three Sample 2 JSON files.
        2. Normalize dimensions and privacy-safe fact tables.
        3. Assess Meta-to-WhatsApp data quality.
        4. Load and validate campaign-type and budget policies.
        5. Aggregate five scorecard levels and calculate KPIs.
        6. Document the derived KPI formulas.
        7. Build campaign evidence packs and benchmarks.
        8. Assign deterministic campaign target and funding decisions.
        9. Propagate decisions to adsets, ads, creatives, and audiences.
        10. Build the illustrative campaign-level budget scenario.
        11. Optionally run the structured LLM narrative pipeline.
        12. Assemble final audit outputs.

        **Known current gaps:** organic/direct outcomes are retained but not surfaced as
        a formal baseline; daily rows are rolled up but trend/fatigue features are not
        calculated; historical benchmarks are not yet available; child entities receive
        actions but not explicit budget units; the test envelope follows historical
        experimental spend instead of a fixed 30% reserve.
        """
    ),
    code(
        """
        # Step 0 - Runtime setup and Excel audit logger
        from __future__ import annotations

        import json
        import os
        import sys
        from datetime import datetime, timezone
        from pathlib import Path

        import pandas as pd
        from IPython.display import display
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        ROOT = Path.cwd().resolve()
        if not (ROOT / "src_2").exists():
            raise RuntimeError("Run this notebook from the Team2-MarketingExpert repository root.")
        if str(ROOT) not in sys.path:
            sys.path.insert(0, str(ROOT))

        from dotenv import load_dotenv
        load_dotenv(ROOT / ".env")

        from src_2.paths import INPUT_DIR

        INPUT_DIRECTORY = Path(INPUT_DIR)
        OUTPUT_DIRECTORY = ROOT / "outputs" / "sample2_pipeline_audit"
        WORKBOOK_PATH = OUTPUT_DIRECTORY / "sample2_pipeline_step_log.xlsx"
        FINAL_JSON_PATH = OUTPUT_DIRECTORY / "sample2_completed_cycle_report.json"
        RUN_LLM = False
        MODEL = os.getenv("OPENAI_MODEL", "gpt-5-mini")

        OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
        if WORKBOOK_PATH.exists():
            WORKBOOK_PATH.unlink()

        SHEET_LOG: dict[str, dict] = {
            "00_Run_Log": {
                "sheet_sequence": 1,
                "worksheet_name": "00_Run_Log",
                "pipeline_step": "00_runtime_setup",
                "data_rows": 0,
                "data_columns": 7,
                "purpose": "Index of every physical worksheet written to this workbook",
                "last_written_utc": None,
            }
        }


        def _excel_value(value):
            if isinstance(value, (dict, list, tuple, set)):
                value = json.dumps(value, ensure_ascii=True, default=str)
            if pd.isna(value) if not isinstance(value, str) else False:
                return None
            text = str(value) if not isinstance(value, (int, float, bool, datetime)) else value
            return text[:32000] if isinstance(text, str) else text


        def _excel_frame(value) -> pd.DataFrame:
            if isinstance(value, pd.Series):
                frame = value.to_frame().reset_index()
            elif isinstance(value, pd.DataFrame):
                frame = value.copy()
            elif isinstance(value, dict):
                frame = pd.DataFrame([value])
            else:
                frame = pd.DataFrame(value)
            frame.columns = [str(column) for column in frame.columns]
            for column in frame.columns:
                if isinstance(frame[column].dtype, pd.DatetimeTZDtype):
                    frame[column] = frame[column].dt.tz_convert(None)
                elif frame[column].dtype == "object":
                    frame[column] = frame[column].map(_excel_value)
            return frame.replace([float("inf"), float("-inf")], None)


        def _style_sheets(sheet_names: list[str]) -> None:
            workbook = load_workbook(WORKBOOK_PATH)
            header_fill = PatternFill("solid", fgColor="17365D")
            header_font = Font(color="FFFFFF", bold=True)
            for sheet_name in sheet_names:
                worksheet = workbook[sheet_name]
                worksheet.sheet_view.showGridLines = False
                worksheet.freeze_panes = "A2"
                if worksheet.max_row >= 1 and worksheet.max_column >= 1:
                    worksheet.auto_filter.ref = worksheet.dimensions
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                for column_index in range(1, worksheet.max_column + 1):
                    header = str(worksheet.cell(1, column_index).value or "").lower()
                    values = [worksheet.cell(row, column_index).value for row in range(1, min(worksheet.max_row, 250) + 1)]
                    width = min(max(len(str(value)) for value in values if value is not None) + 2, 42) if any(value is not None for value in values) else 12
                    worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)
                    for row in range(2, worksheet.max_row + 1):
                        cell = worksheet.cell(row, column_index)
                        cell.alignment = Alignment(vertical="top", wrap_text=False)
                        if any(token in header for token in ("date", "timestamp", "extracted_at")):
                            cell.number_format = "yyyy-mm-dd hh:mm"
                        elif any(token in header for token in ("spend", "revenue", "value", "budget_units", "benchmark", "actual")):
                            cell.number_format = "#,##0.00"
                        elif header.endswith("_rate") or header.endswith("_ratio"):
                            cell.number_format = "0.00%"
                worksheet.row_dimensions[1].height = 30
            workbook.save(WORKBOOK_PATH)


        def _sheet_purpose(sheet_name: str) -> str:
            readable = sheet_name.split("_", 1)[-1].replace("_", " ").lower()
            return f"Pipeline output table: {readable}"


        def log_step(step: str, sheets: dict[str, pd.DataFrame], note: str = "") -> None:
            prepared = {name[:31]: _excel_frame(frame) for name, frame in sheets.items()}
            timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
            for sheet_name, frame in prepared.items():
                existing = SHEET_LOG.get(sheet_name)
                SHEET_LOG[sheet_name] = {
                    "sheet_sequence": existing["sheet_sequence"] if existing else len(SHEET_LOG) + 1,
                    "worksheet_name": sheet_name,
                    "pipeline_step": step,
                    "data_rows": len(frame),
                    "data_columns": len(frame.columns),
                    "purpose": _sheet_purpose(sheet_name),
                    "last_written_utc": timestamp,
                }
            SHEET_LOG["00_Run_Log"].update(
                {
                    "data_rows": len(SHEET_LOG),
                    "last_written_utc": timestamp,
                }
            )
            run_log = pd.DataFrame(SHEET_LOG.values()).sort_values("sheet_sequence")
            frames = {"00_Run_Log": run_log, **prepared}
            mode = "a" if WORKBOOK_PATH.exists() else "w"
            options = {"engine": "openpyxl", "mode": mode}
            if mode == "a":
                options["if_sheet_exists"] = "replace"
            with pd.ExcelWriter(WORKBOOK_PATH, **options) as writer:
                for sheet_name, frame in frames.items():
                    frame.to_excel(writer, sheet_name=sheet_name, index=False)
            _style_sheets(list(frames))
            print(
                f"Logged {step}: {len(prepared)} worksheet(s); "
                f"{len(SHEET_LOG)} physical worksheets recorded"
            )


        guide = pd.DataFrame(
            [
                {"setting": "input_directory", "value": str(INPUT_DIRECTORY), "meaning": "Folder containing the three Sample 2 JSON files"},
                {"setting": "workbook", "value": str(WORKBOOK_PATH), "meaning": "Audit workbook updated after every completed step"},
                {"setting": "run_llm", "value": RUN_LLM, "meaning": "False avoids the 14 paid narrative calls"},
                {"setting": "model", "value": MODEL, "meaning": "OpenAI model used when RUN_LLM is True"},
            ]
        )
        log_step("00_runtime_setup", {"00_Run_Guide": guide})
        display(guide)
        """
    ),
    markdown(
        """
        ## Step 1 - Load the raw cycle

        The loader checks that all three required files exist and have the expected
        top-level JSON type. The inventory check confirms that the Meta object contains
        campaign, adset, ad, creative, and insight collections.
        """
    ),
    code(
        """
        from src_2.ingestion import load_sample2
        from src_2.ingestion.data_quality import inventory

        raw_payload = load_sample2(INPUT_DIRECTORY)
        source_inventory = inventory(raw_payload)
        inventory_frame = pd.DataFrame(
            [{"collection": key, "record_count": value} for key, value in vars(source_inventory).items()]
        )
        source_files = pd.DataFrame(
            [
                {"file": name, "path": str(INPUT_DIRECTORY / name), "size_bytes": (INPUT_DIRECTORY / name).stat().st_size}
                for name in ("meta_data.json", "conversations.json", "products.json")
            ]
        )
        log_step(
            "01_load_raw_json",
            {"01_Source_Inventory": inventory_frame, "01_Source_Files": source_files},
            "No business calculations are performed in this step.",
        )
        display(inventory_frame)
        """
    ),
    markdown(
        """
        ## Step 2 - Normalize canonical dimensions and facts

        IDs are standardized as strings, dates become timestamps, numeric media fields
        become numbers, and duplicates are removed. WhatsApp rows are linked to the Meta
        hierarchy through source IDs and the ad lookup.

        Organic and direct conversations remain in the canonical conversation table, but
        have no paid campaign ID. They will therefore be excluded from paid scorecards.
        Raw message text, phone numbers, and customer names are not copied into canonical
        data or Excel.
        """
    ),
    code(
        """
        from src_2.ingestion import normalize_cycle

        canonical = normalize_cycle(raw_payload)
        source_context = (
            canonical.conversations.groupby("source_platform", dropna=False)
            .agg(
                conversations=("conversation_id", "nunique"),
                unique_customers=("customer_id", "nunique"),
                delivered_orders=("is_delivered", "sum"),
                net_revenue=("net_revenue", "sum"),
            )
            .reset_index()
        )
        canonical_inventory = pd.DataFrame(
            [
                {"canonical_table": name, "rows": len(getattr(canonical, name)), "natural_grain": grain}
                for name, grain in {
                    "campaigns": "one row per campaign",
                    "adsets": "one row per adset",
                    "ads": "one row per ad",
                    "creatives": "one row per creative",
                    "media_daily": "one row per ad and date",
                    "conversations": "one row per conversation",
                    "order_lines": "one row per ordered product line",
                    "products": "one row per product",
                }.items()
            ]
        )
        log_step(
            "02_normalize_canonical_data",
            {
                "02_Canonical_Inventory": canonical_inventory,
                "02_Source_Context": source_context,
                "02_Campaigns": canonical.campaigns,
                "02_Adsets": canonical.adsets,
                "02_Ads": canonical.ads,
                "02_Creatives": canonical.creatives,
                "02_Media_Daily": canonical.media_daily,
                "02_Conversations": canonical.conversations,
                "02_Order_Lines": canonical.order_lines,
                "02_Products": canonical.products,
            },
            "Organic/direct are retained for context and excluded later when entity IDs are missing.",
        )
        display(canonical_inventory)
        display(source_context)
        """
    ),
    markdown(
        """
        ## Step 3 - Assess evidence quality

        This compares Meta-attributed conversation starts with supplied Meta-sourced
        WhatsApp conversations and checks whether their campaign/adset/ad IDs resolve.
        The ratio is a reconciliation warning, not proven literal coverage, because the
        two sources may use different event definitions, windows, and deduplication.
        """
    ),
    code(
        """
        from src_2.ingestion import build_data_quality_report

        data_quality = build_data_quality_report(canonical)
        quality_frame = pd.DataFrame([data_quality.model_dump(mode="json")]).drop(columns=["warnings"])
        quality_warnings = pd.DataFrame({"warning": data_quality.warnings})
        log_step(
            "03_build_data_quality_report",
            {"03_Data_Quality": quality_frame, "03_Quality_Warnings": quality_warnings},
            "The current Sample 2 evidence status is limited_evidence.",
        )
        display(quality_frame)
        display(quality_warnings)
        """
    ),
    markdown(
        """
        ## Step 4 - Load business rules

        Campaign-type YAML selects the business job, primary KPI, supporting KPIs,
        guardrails, and one allocation metric. Budget YAML defines the illustrative
        100-unit allocation policy. Pydantic rejects missing, invalid, or extra fields.
        """
    ),
    code(
        """
        from src_2.infrastructure.configuration import load_budget_policy, load_campaign_type_registry

        campaign_registry = load_campaign_type_registry()
        budget_policy = load_budget_policy()
        campaign_rules = []
        for campaign_type, rule in campaign_registry.campaign_types.items():
            campaign_rules.append(
                {
                    "campaign_type": campaign_type.value,
                    "business_job": rule.business_job,
                    "success_question": rule.success_question,
                    "primary_kpis": ", ".join(rule.primary_kpis),
                    "supporting_kpis": ", ".join(rule.supporting_kpis),
                    "allocation_metric": rule.allocation_metric.metric,
                    "allocation_direction": rule.allocation_metric.direction,
                    "budget_pool": rule.budget_pool.value,
                    "guardrails": [guard.model_dump(mode="json") for guard in rule.guardrails],
                }
            )
        campaign_rules_frame = pd.DataFrame(campaign_rules)
        budget_policy_frame = pd.json_normalize(budget_policy.model_dump(mode="json"), sep=".")
        action_eligibility = pd.DataFrame(
            [
                {"budget_pool": pool, "eligible_action": action}
                for pool, actions in budget_policy.action_eligibility.items()
                for action in actions
            ]
        )
        log_step(
            "04_load_business_configuration",
            {
                "04_Campaign_Rules": campaign_rules_frame,
                "04_Budget_Policy": budget_policy_frame,
                "04_Action_Eligibility": action_eligibility,
            },
            "The current test envelope follows historical experimental spend, not a fixed 30% reserve.",
        )
        display(campaign_rules_frame)
        """
    ),
    markdown(
        """
        ## Step 5 - Build five scorecards

        Media, conversations, and product lines are aggregated independently at each
        entity level and only then merged. Ratios are recomputed from aggregated totals;
        daily ratios are not averaged.
        """
    ),
    code(
        """
        from src_2.analytics import build_scorecards

        raw_scorecards = build_scorecards(canonical)
        reconciliation_rows = []
        for level in ("campaign", "adset", "ad", "creative", "audience"):
            frame = raw_scorecards.by_level(level)
            reconciliation_rows.append(
                {
                    "level": level,
                    "entities": len(frame),
                    "spend": frame["spend"].sum(),
                    "observed_conversations": frame["observed_conversations"].sum(),
                    "net_revenue": frame["net_revenue"].sum(),
                }
            )
        reconciliation = pd.DataFrame(reconciliation_rows)
        log_step(
            "05_build_scorecards",
            {
                "05_Reconciliation": reconciliation,
                "05_Campaign_Scorecard": raw_scorecards.campaign,
                "05_Adset_Scorecard": raw_scorecards.adset,
                "05_Ad_Scorecard": raw_scorecards.ad,
                "05_Creative_Scorecard": raw_scorecards.creative,
                "05_Audience_Scorecard": raw_scorecards.audience,
            },
            "Spend and attributed outcomes reconcile across every scorecard level.",
        )
        display(reconciliation)
        display(raw_scorecards.campaign.head())
        """
    ),
    markdown(
        """
        ## Step 6 - Document calculated fields

        These definitions explain the calculated columns already present in every
        scorecard. `net_revenue` and `net_roas` are contribution proxies, not profit,
        because product cost, fulfillment, tax, and overhead are unavailable.
        """
    ),
    code(
        """
        kpi_definitions = pd.DataFrame(
            [
                {"metric": "frequency", "calculation": "impressions / reach", "business_meaning": "Average exposure per reached person"},
                {"metric": "link_ctr_pct", "calculation": "link_clicks / impressions * 100", "business_meaning": "Share of impressions producing a link click"},
                {"metric": "cpm", "calculation": "spend / impressions * 1000", "business_meaning": "Cost per 1,000 impressions"},
                {"metric": "cpc", "calculation": "spend / link_clicks", "business_meaning": "Cost per link click"},
                {"metric": "cost_per_observed_conversation", "calculation": "spend / observed WhatsApp conversations", "business_meaning": "Media cost per supplied attributed conversation"},
                {"metric": "cost_per_delivered_order", "calculation": "spend / delivered orders", "business_meaning": "Media cost per delivered order"},
                {"metric": "net_roas", "calculation": "net revenue / spend", "business_meaning": "Observed revenue return per media unit spent"},
                {"metric": "aov", "calculation": "delivered revenue / delivered orders", "business_meaning": "Average delivered order value"},
                {"metric": "delivered_rate", "calculation": "delivered orders / observed conversations", "business_meaning": "Share of supplied conversations becoming delivered orders"},
                {"metric": "negative_outcome_rate", "calculation": "ghosted + cancelled + refunded + adversarial / observed conversations", "business_meaning": "Share of observed conversations ending negatively"},
                {"metric": "net_revenue_per_day", "calculation": "net revenue / active media days", "business_meaning": "Revenue normalized for different campaign durations"},
                {"metric": "repeat_order_rate", "calculation": "repeat delivered orders / delivered orders", "business_meaning": "Delivered-order share attributed to repeat cycles"},
            ]
        )
        log_step(
            "06_document_kpi_definitions",
            {"06_KPI_Definitions": kpi_definitions},
            "All ratios use aggregated numerators and denominators.",
        )
        display(kpi_definitions)
        """
    ),
    markdown(
        """
        ## Step 7 - Build campaign evidence packs

        Each pack is the exact handoff from measurement to the decision and campaign
        analysis stages. It contains campaign-type context, KPI evidence, benchmarks,
        guardrails, limitations, and child-entity evidence.

        The current benchmark resolver excludes the entity being assessed, then uses a
        current-cycle same-type median and falls back to a current-cycle portfolio median.
        Historical observations and approved campaign targets are not yet loaded.
        """
    ),
    code(
        """
        from src_2.analytics import build_evidence_packs
        from src_2.application.reporting import _manifest

        manifest = _manifest(canonical)
        evidence_packs = build_evidence_packs(
            manifest.cycle_id, raw_scorecards, campaign_registry, data_quality
        )
        pack_summary = []
        metric_evidence_rows = []
        child_evidence_rows = []
        for pack in evidence_packs:
            pack_summary.append(
                {
                    "campaign_id": pack.campaign_id,
                    "campaign_name": pack.campaign_name,
                    "campaign_type": pack.campaign_type.value,
                    "business_job": pack.business_job,
                    "success_question": pack.success_question,
                    "evidence_status": pack.evidence_status.value,
                    "adsets": len(pack.adsets),
                    "ads": len(pack.ads),
                    "creatives": len(pack.creatives),
                    "audiences": len(pack.audiences),
                }
            )
            for category, metrics in (
                ("primary", pack.primary_kpis),
                ("supporting", pack.supporting_kpis),
                ("guardrail", pack.guardrails),
                ("allocation", [pack.allocation_kpi] if pack.allocation_kpi else []),
            ):
                for metric in metrics:
                    metric_evidence_rows.append(
                        {"campaign_id": pack.campaign_id, "category": category, **metric.model_dump(mode="json")}
                    )
            for entities in (pack.adsets, pack.ads, pack.creatives, pack.audiences):
                for entity in entities:
                    for metric in entity.metrics:
                        child_evidence_rows.append(
                            {
                                "campaign_id": pack.campaign_id,
                                "entity_level": entity.entity_level.value,
                                "entity_id": entity.entity_id,
                                "entity_name": entity.entity_name,
                                "metric": metric.metric,
                                "actual": metric.actual,
                                "benchmark": metric.benchmark,
                                "passed": metric.passed,
                                "evidence_count": metric.evidence_count,
                            }
                        )
        pack_summary_frame = pd.DataFrame(pack_summary)
        metric_evidence_frame = pd.DataFrame(metric_evidence_rows)
        child_evidence_frame = pd.DataFrame(child_evidence_rows)
        log_step(
            "07_build_evidence_packs",
            {
                "07_Evidence_Packs": pack_summary_frame,
                "07_Campaign_Evidence": metric_evidence_frame,
                "07_Child_Evidence": child_evidence_frame,
            },
            "One typed CampaignEvidencePack is created for each campaign.",
        )
        display(pack_summary_frame)
        display(metric_evidence_frame.head(20))
        """
    ),
    markdown(
        """
        ## Step 8 - Assign campaign target and funding decisions

        The deterministic assessor answers two different questions:

        - `target_status`: did the campaign perform its campaign-type business job?
        - `next_cycle_action`: should the campaign scale, remain a test, receive no
          funding, or wait for more evidence?

        Limited evidence prevents an automatic scale action.
        """
    ),
    code(
        """
        from src_2.analytics import DeterministicCampaignAssessor

        assessor = DeterministicCampaignAssessor()
        assessments = [
            assessor.assess(pack, campaign_registry.campaign_types[pack.campaign_type])
            for pack in evidence_packs
        ]
        assessment_frame = pd.DataFrame(
            [
                {
                    "campaign_id": item.campaign_id,
                    "campaign_name": item.campaign_name,
                    "campaign_type": item.campaign_type.value,
                    "target_status": item.target_status.value,
                    "next_cycle_action": item.next_cycle_action.value,
                    "evidence_status": item.evidence_status.value,
                    "reason_codes": item.reason_codes,
                }
                for item in assessments
            ]
        )
        assessment_details = pd.DataFrame(
            [
                {
                    "campaign_id": item.campaign_id,
                    "result_type": result_type,
                    **metric.model_dump(mode="json"),
                }
                for item in assessments
                for result_type, results in (("primary", item.primary_results), ("guardrail", item.guardrail_results))
                for metric in results
            ]
        )
        log_step(
            "08_assess_campaigns",
            {"08_Campaign_Decisions": assessment_frame, "08_Decision_Details": assessment_details},
            "Funding actions are deterministic and independent from the LLM narrative.",
        )
        display(assessment_frame)
        """
    ),
    markdown(
        """
        ## Step 9 - Add child-entity decisions

        Adsets, ads, creatives, and audiences are compared with peers inside their parent
        campaign using the campaign type's allocation metric. A blocked parent blocks its
        children. Minimum observed-conversation thresholds are 5 for adsets/audiences and
        3 for ads/creatives.
        """
    ),
    code(
        """
        from src_2.analytics import enrich_scorecards

        scorecards = enrich_scorecards(
            raw_scorecards, assessments, campaign_registry, data_quality
        )
        child_columns = [
            "campaign_id", "campaign_name", "entity_id", "entity_name",
            "spend", "observed_conversations", "delivered_orders", "net_revenue",
            "allocation_metric", "allocation_metric_value", "allocation_benchmark",
            "allocation_metric_passed", "evidence_status", "next_cycle_action", "decision_reason",
        ]
        child_decision_sheets = {}
        child_summary_rows = []
        for number, level in enumerate(("adset", "ad", "creative", "audience"), start=1):
            frame = scorecards.by_level(level)
            selected = frame[[column for column in child_columns if column in frame]].copy()
            child_decision_sheets[f"09_{level.title()}_Decisions"] = selected
            for action, count in selected["next_cycle_action"].value_counts().items():
                child_summary_rows.append({"entity_level": level, "next_cycle_action": action, "entities": count})
        child_summary = pd.DataFrame(child_summary_rows)
        child_decision_sheets["09_Child_Summary"] = child_summary
        log_step(
            "09_enrich_child_decisions",
            child_decision_sheets,
            "Children receive actions and reasons, but the current allocator does not assign child budget units.",
        )
        display(child_summary)
        """
    ),
    markdown(
        """
        ## Step 10 - Build the normalized budget scenario

        The allocator starts with 100 illustrative units, preserves each campaign type's
        historical spend share, filters campaigns by deterministic action eligibility,
        and distributes each type envelope using one configured allocation KPI.

        Money blocked by the rules remains unallocated. This output is not an approved
        currency budget and is non-operational while evidence is limited.
        """
    ),
    code(
        """
        from src_2.analytics import DeterministicBudgetAllocator

        allocator = DeterministicBudgetAllocator()
        budget_scenario = allocator.allocate(
            manifest.cycle_id,
            scorecards.campaign,
            assessments,
            campaign_registry,
            budget_policy,
            data_quality,
        )
        budget_summary = pd.DataFrame(
            [
                {
                    "cycle_id": budget_scenario.cycle_id,
                    "scenario_name": budget_scenario.scenario_name,
                    "total_budget_units": budget_scenario.total_budget_units,
                    "allocated_units": sum(item.budget_units for item in budget_scenario.allocations),
                    "unallocated_units": budget_scenario.unallocated_units,
                    "evidence_status": budget_scenario.evidence_status.value,
                    "operational": budget_scenario.operational,
                }
            ]
        )
        budget_allocations = pd.DataFrame(
            [
                {
                    **item.model_dump(mode="json"),
                    "entity_level": item.entity_level.value,
                    "action": item.action.value,
                }
                for item in budget_scenario.allocations
            ]
        ).sort_values("budget_units", ascending=False)
        budget_assumptions = pd.DataFrame({"assumption": budget_scenario.assumptions})
        log_step(
            "10_allocate_normalized_budget",
            {
                "10_Budget_Summary": budget_summary,
                "10_Budget_Allocations": budget_allocations,
                "10_Budget_Assumptions": budget_assumptions,
            },
            "The current Sample 2 scenario remains illustrative and non-operational.",
        )
        display(budget_summary)
        display(budget_allocations[["entity_name", "action", "budget_units", "reason"]])
        """
    ),
    markdown(
        """
        ## Step 11 - Optional structured LLM narrative

        When enabled, the notebook performs:

        1. One campaign-analysis call per campaign: `CampaignEvidencePack` -> `CampaignInsight`.
        2. One portfolio call: assessments plus campaign insights -> `PortfolioInsight`.
        3. One narrator call: portfolio insight plus fixed budget -> `StakeholderReport`.

        Pydantic validates every response. The LLM explains supplied evidence and fixed
        decisions; it does not recalculate KPIs, target statuses, actions, or budget.
        """
    ),
    code(
        """
        from src_2.intelligence import OpenAICampaignAnalyst, OpenAIPortfolioSynthesizer, OpenAIReportNarrator

        llm_input_summary = pd.DataFrame(
            [
                {
                    "campaign_id": pack.campaign_id,
                    "campaign_name": pack.campaign_name,
                    "campaign_type": pack.campaign_type.value,
                    "evidence_status": pack.evidence_status.value,
                    "primary_kpis": ", ".join(metric.metric for metric in pack.primary_kpis),
                    "guardrails": ", ".join(metric.metric for metric in pack.guardrails),
                    "adsets": len(pack.adsets),
                    "ads": len(pack.ads),
                    "creatives": len(pack.creatives),
                    "audiences": len(pack.audiences),
                }
                for pack in evidence_packs
            ]
        )

        campaign_insights = []
        portfolio_insight = None
        stakeholder_report = None
        llm_sheets = {"11_LLM_Input_Summary": llm_input_summary}

        if RUN_LLM:
            if not os.getenv("OPENAI_API_KEY"):
                raise RuntimeError("RUN_LLM is True but OPENAI_API_KEY is missing from .env")
            analyst = OpenAICampaignAnalyst(model=MODEL)
            campaign_insights = [analyst.analyze(pack) for pack in evidence_packs]
            portfolio_insight = OpenAIPortfolioSynthesizer(model=MODEL).synthesize(
                assessments, campaign_insights
            )
            stakeholder_report = OpenAIReportNarrator(model=MODEL).narrate(
                portfolio_insight, budget_scenario
            )
            llm_status = pd.DataFrame([{"run_llm": True, "model": MODEL, "api_calls": len(evidence_packs) + 2, "status": "completed"}])
            llm_sheets["11_Campaign_Insights"] = pd.DataFrame(
                [item.model_dump(mode="json") for item in campaign_insights]
            )
            llm_sheets["11_Portfolio_Insight"] = pd.DataFrame(
                [portfolio_insight.model_dump(mode="json")]
            )
            llm_sheets["11_Stakeholder_Report"] = pd.DataFrame(
                [stakeholder_report.model_dump(mode="json")]
            )
        else:
            llm_status = pd.DataFrame(
                [{"run_llm": False, "model": MODEL, "api_calls": 0, "status": "skipped", "reason": "Set RUN_LLM=True to execute the 14-call narrative pipeline."}]
            )
        llm_sheets["11_LLM_Status"] = llm_status
        log_step(
            "11_optional_llm_narrative",
            llm_sheets,
            "LLM execution is optional; all KPI, decision, and budget outputs already exist deterministically.",
        )
        display(llm_status)
        """
    ),
    markdown(
        """
        ## Step 12 - Assemble final outputs

        The audit workbook is always complete through the deterministic budget stage.
        When the LLM section is enabled, this cell also assembles the same
        `CompletedCycleReport` used by Streamlit and exports its structured JSON.
        """
    ),
    code(
        """
        from src_2.application.reporting import CompletedCycleReport

        report = None
        final_summary = pd.DataFrame(
            [
                {
                    "cycle_id": manifest.cycle_id,
                    "reporting_start": manifest.reporting_start,
                    "reporting_end": manifest.reporting_end,
                    "campaigns": len(scorecards.campaign),
                    "campaigns_keep_as_test": sum(item.next_cycle_action.value == "keep_as_test" for item in assessments),
                    "campaigns_do_not_fund": sum(item.next_cycle_action.value == "do_not_fund" for item in assessments),
                    "allocated_budget_units": sum(item.budget_units for item in budget_scenario.allocations),
                    "unallocated_budget_units": budget_scenario.unallocated_units,
                    "operational": budget_scenario.operational,
                    "llm_report_generated": stakeholder_report is not None,
                    "excel_sheets_logged": len(SHEET_LOG) + 1,
                    "excel_workbook": str(WORKBOOK_PATH),
                }
            ]
        )

        final_sheets = {"12_Final_Summary": final_summary}
        if stakeholder_report is not None:
            report = CompletedCycleReport(
                manifest=manifest,
                data_quality=data_quality,
                canonical_data=canonical,
                scorecards=scorecards,
                evidence_packs=evidence_packs,
                assessments=assessments,
                insights=campaign_insights,
                portfolio_insight=portfolio_insight,
                budget_scenario=budget_scenario,
                stakeholder_report=stakeholder_report,
            )
            FINAL_JSON_PATH.write_text(
                json.dumps(report.to_export_dict(), indent=2, ensure_ascii=True),
                encoding="utf-8",
            )
            final_sheets["12_Final_Artifacts"] = pd.DataFrame(
                [{"artifact": "completed_cycle_report_json", "path": str(FINAL_JSON_PATH)}]
            )

        log_step(
            "12_finalize_outputs",
            final_sheets,
            "The Excel workbook contains an auditable snapshot of every completed stage.",
        )
        workbook_sheet_names = load_workbook(WORKBOOK_PATH, read_only=True).sheetnames
        actual_sheet_log = pd.DataFrame(SHEET_LOG.values()).sort_values("sheet_sequence")
        assert len(workbook_sheet_names) == len(actual_sheet_log)
        assert set(workbook_sheet_names) == set(actual_sheet_log["worksheet_name"])
        display(final_summary)
        display(actual_sheet_log)
        print(f"Audit workbook: {WORKBOOK_PATH}")
        print(f"Actual worksheets logged: {len(actual_sheet_log)}")
        if report is None:
            print("Structured report JSON was not generated because RUN_LLM=False.")
        else:
            print(f"Structured report JSON: {FINAL_JSON_PATH}")
        """
    ),
    markdown(
        """
        ## How to use the workbook

        Start with `00_Run_Log`, then follow sheets by numeric prefix. The most important
        business sheets are:

        - `03_Data_Quality`: whether recommendations can be operational.
        - `05_Campaign_Scorecard`: campaign economics and outcomes.
        - `07_Campaign_Evidence`: benchmarks and pass/fail evidence.
        - `08_Campaign_Decisions`: target achievement and next-cycle action.
        - `09_*_Decisions`: adset, ad, creative, and audience actions.
        - `10_Budget_Allocations`: illustrative next-cycle campaign budget.
        - `11_*`: optional LLM input and output audit.

        The workbook is a trace of Python-calculated outputs. Business assumptions remain
        visible in YAML and are copied into the configuration and budget sheets.
        """
    ),
]


notebook = nbf.v4.new_notebook(cells=cells)
notebook.metadata = {
    "kernelspec": {
        "display_name": "Python 3",
        "language": "python",
        "name": "python3",
    },
    "language_info": {"name": "python", "version": "3.9"},
}
nbf.write(notebook, OUTPUT)
print(OUTPUT)
