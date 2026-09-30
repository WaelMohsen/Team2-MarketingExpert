from pathlib import Path
import textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT

BASE = Path(__file__).resolve().parent
OUT = BASE.parents[1] / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
maps = [
('Function overview', 'score_level', [
('1  Prepare inputs', 'frame: entities at one level; level: entity type; registry: objective rules. Copy frame and create scored list.'),
('2–3  Read objective and evidence', 'For each row: load its contract, read successes and trials, and classify evidence sufficiency.'),
('4–5  Choose comparison and initialize', 'Prefer same-objective peers; try a compatible fallback. Set raw metrics, comparison metadata and default values.'),
('6–7  Assess performance and efficiency', 'Estimate an adjusted rate and uncertainty, derive a statistical decision, then compare cost or return with peers.'),
('8–9  Recommend and return', 'Apply action rules, save reason codes, repeat for every row, and merge scores into the original data.')],
'Scope: score_level produces primary scores and recommendations. Semantic scoring and budget allocation run later.'),
('Steps 1 to 3 Inputs and evidence', 'Read each row', [
('Prepare', 'output = frame.copy(); scored = []. Iterate through output rows without changing the input frame.'),
('Validate objective', 'Read row.objective. If absent from registry.objectives, raise ValueError; otherwise load its contract.'),
('Read counts', 'successes uses primary_numerator; trials uses primary_denominator. Convert to float; absent fields default to zero.'),
('Classify evidence', 'trials <= 0: none. Below minimum_trials: limited. At or above the minimum: sufficient. Current minimum: 10.'),
('Interpret the denominator', 'Awareness / Engagement: impressions. Leads: mature conversations. Sales: mature unique customers.')],
'Evidence classification does not stop calculations. The final action checks sufficiency again.'),
('Step 4 Select the benchmark', 'Find valid peers', [
('Filter peers', 'Same level and objective; exclude current entity_id. Require trials > 0 and 0 <= successes <= trials.'),
('Enough same-objective peers', 'At least minimum_peer_entities (currently 2): scope = same_objective; quality = decision_grade.'),
('Try a compatible fallback', 'Require the same fallback group, primary metric, numerator, denominator and direction. Current option: Awareness + Engagement Link CTR.'),
('Evaluate fallback', 'Enough valid fallback peers: scope = shared_primary_kpi_group; quality = provisional. Otherwise comparison remains unavailable.'),
('Portfolio context', 'Compute an adjusted pooled rate across valid other entities using the selected count columns. Descriptive only; it does not decide the action.')],
'Provisional comparisons can only support KEEP_AS_TEST when evidence is sufficient. The current entity never benchmarks itself.'),
('Step 5 Initialize the score record', 'Build values', [
('Identity and evidence', 'entity_id; score_metric; score_numerator; score_denominator; score_direction; score_successes; score_trials; evidence_status.'),
('Raw rate', 'raw_rate = successes / trials when trials is nonzero; otherwise None.'),
('Benchmark metadata', 'benchmark_peer_count; same_objective_peer_count; benchmark_scope; benchmark_quality; portfolio_context_benchmark and peer count.'),
('Empty estimates', 'Initialize corrected_rate, rate bounds, benchmark, expected_lift, lift bounds and probability_better to None.'),
('Default decision and efficiency settings', 'statistical_decision = insufficient_evidence. Store efficiency_metric and efficiency_direction from the objective contract.')],
'These defaults remain when statistical estimation cannot run, avoiding unsupported numerical estimates.'),
('Step 6 Build the statistical estimate', 'Estimate rates', [
('Check calculation conditions', 'Run only if trials > 0, benchmark quality is available, and the peer count meets the configured minimum.'),
('Fit the peer prior', 'Mean = (peer successes + 0.5) / (peer trials + 1). Start strength at average peer trials; reduce it when peer-rate variation warrants it; minimum strength is 1.'),
('Construct Beta parameters', 'alpha = mean × strength; beta = (1 - mean) × strength. The helper applies a tiny positive floor to each parameter.'),
('Update with entity evidence', 'Posterior alpha = alpha + successes. Posterior beta = beta + trials - successes. Corrected rate = posterior alpha / sum of posterior parameters.'),
('Simulate uncertainty', 'Use 5,000 posterior draws and 5,000 prior draws. A stable seed uses level, entity_id and metric. Rate bounds are posterior percentiles 2.5 and 97.5.')],
'The calculation can run below 10 trials, but limited evidence still blocks a final funding recommendation.'),
('Step 6 Derive the statistical decision', 'Compare draws', [
('Define favorable lift', 'For higher-is-better metrics: entity draw minus benchmark draw. For lower-is-better metrics: benchmark draw minus entity draw.'),
('Summarize comparisons', 'expected_lift = mean difference. probability_better = fraction of positive differences. benchmark = prior mean.'),
('Find the lift interval', 'lift_low and lift_high are the 2.5th and 97.5th percentiles of the paired differences.'),
('Classify performance', 'lift_low > 0: scale. Otherwise, lift_high < 0: kill. Otherwise: hold, including intervals that touch zero.'),
('Keep the distinction', 'The statistical decision is an intermediate result. A high probability_better alone does not guarantee SCALE.')],
'Example: a lift interval from -4 to +12 percentage points produces hold, even when the average lift is positive.'),
('Step 7 Compare efficiency', 'Check cost or return', [
('Select comparison values', 'Use other entities with the same objective. Convert the configured efficiency column to numbers and drop missing values.'),
('Handle missing evidence', 'If the entity value is missing or no peer values remain, comparison = unavailable and benchmark = None.'),
('Compute the benchmark', 'Use the median peer value, not the mean. One available peer value can support this comparison; there is no two-peer gate here.'),
('Respect direction', 'Lower is better for CPM, CPC and cost per created order. Higher is better for Net ROAS.'),
('Return efficiency fields', 'Classify approximately equal values as equal; otherwise better or worse. Save value, median benchmark, peer count and comparison.')],
'Efficiency always uses same-objective peers. It gates final SCALE recommendations rather than changing the primary statistical score.'),
('Step 8 Apply the final action rules', 'First matching rule', [
('1  Insufficient evidence', 'Evidence is not sufficient OR statistical_decision is insufficient_evidence → INSUFFICIENT_EVIDENCE.'),
('2  Provisional comparison', 'If the first rule did not apply and benchmark quality is provisional → KEEP_AS_TEST, regardless of scale, hold or kill.'),
('3  Clear underperformance', 'With sufficient evidence and a decision-grade comparison, statistical kill → DO_NOT_FUND.'),
('4  Inconclusive comparison', 'With sufficient evidence and a decision-grade comparison, statistical hold → KEEP_AS_TEST.'),
('5  Clear outperformance', 'For statistical scale: efficiency better/equal → SCALE. Efficiency worse/unavailable → KEEP_AS_TEST.')],
'Rule order matters. Each outcome includes reason codes explaining evidence, lift, fallback comparison or efficiency.'),
('Step 9 Save and return results', 'Complete the loop', [
('Receive the action', '_action returns an action enum and a list of reasons.'),
('Save the recommendation', 'values[recommended_action] = action.value; values[reason_codes] = reasons.'),
('Append and repeat', 'Append values to scored. If another row exists, repeat from objective lookup.'),
('Merge the score records', 'Convert scored into a DataFrame. Left-merge it with output using entity_id.'),
('Return the enriched table', 'Preserve original entity data and add primary estimates, evidence and benchmark details, efficiency, recommendation and reasons.')],
'Source: src_mvp/statistics.py, score_level and its helpers; objective defaults: src_mvp/config/objectives.yaml.')
]

def diagram(index, root, branches):
    im = Image.new('RGB', (2400, 1330), 'white')
    d = ImageDraw.Draw(im)
    titlefont = ImageFont.truetype(BOLD, 34)
    bodyfont = ImageFont.truetype(FONT, 32)
    rootfont = ImageFont.truetype(BOLD, 40)
    center = 660
    d.rounded_rectangle((35, center-95, 485, center+95), 20, fill='#17384F')
    lines = textwrap.wrap(root, 18)
    for j, line in enumerate(lines):
        d.text((260, center+(j-(len(lines)-1)/2)*49), line, font=rootfont, fill='white', anchor='mm')
    d.line((485,center,560,center), fill='#8193A0', width=4)
    ys = [130 + n*260 for n in range(len(branches))]
    d.line((560,ys[0],560,ys[-1]), fill='#8193A0', width=4)
    for y, (heading, body) in zip(ys, branches):
        d.line((560,y,630,y), fill='#8193A0', width=4)
        d.rounded_rectangle((630,y-112,2370,y+112), 14, fill='#F0F5F8', outline='#CCD8DF', width=2)
        d.text((662,y-87), heading, font=titlefont, fill='#17384F')
        wrapped = textwrap.wrap(body, width=98)
        assert len(wrapped) <= 4, (heading, wrapped)
        for j,line in enumerate(wrapped):
            d.text((662,y-34+j*39), line, font=bodyfont, fill='#26333C')
    path = BASE / f'map-{index}.png'
    im.save(path)
    return path

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Inches(11.7); sec.page_height=Inches(8.3)
sec.top_margin=sec.bottom_margin=Inches(.5)
sec.left_margin=sec.right_margin=Inches(.6)
for name in ['Normal','Title','Heading 1']:
    doc.styles[name].font.name='Arial'
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(21)
for i,(title,root,branches,note) in enumerate(maps,1):
    if i>1: doc.add_page_break()
    doc.add_paragraph('Score Level Decision Maps' if i==1 else title, 'Title' if i==1 else 'Heading 1')
    doc.add_paragraph('English guide to the function and its helper calls' if i==1 else f'Map {i} of {len(maps)}  |  score_level')
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
    run=p.add_run(); run.add_picture(str(diagram(i,root,branches)), width=Inches(10.45))
    inline=run._r.xpath('.//wp:docPr')[0]
    inline.set('descr', title + '. ' + ' '.join(h+': '+b for h,b in branches))
    doc.add_paragraph(note)
doc.core_properties.title='Score Level Decision Maps'
doc.core_properties.subject='English mind maps of score_level execution and decision rules'
doc.save(OUT/'Score_Level_Decision_Maps.docx')
print(OUT/'Score_Level_Decision_Maps.docx')
